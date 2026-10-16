"""Проверка публичных выгрузок и сборка меню; исходных рецептур здесь нет."""
import argparse
import csv
import hashlib
import ipaddress
import os
import socket
import sys
import tempfile
from decimal import Decimal, ROUND_HALF_UP, getcontext
from pathlib import Path

getcontext().prec = 50
ROOT = Path(__file__).resolve().parents[1]
D = Decimal
ARTIFACTS = ("results/menu_preview.png", "results/cost_50ml.png", "results/nastoyka_menu.pdf")
BASIC = {"C01", "C02", "C03", "C04", "C05", "C09", "C13", "C14"}
PREMIUM = {"C06", "C12"}


def read_csv(name):
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def number(value):
    return D(str(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))


def fmt(value, digits=2):
    value = number(value).quantize(D(1).scaleb(-digits), rounding=ROUND_HALF_UP)
    return f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")


def research_summary():
    """Повторить прежние определения срезов из исходного Excel, не меняя книгу."""
    from openpyxl import load_workbook
    with (ROOT / "data/nastoyki_reviews.xlsx").open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        def rows(sheet):
            values = workbook[sheet].iter_rows(values_only=True)
            header = next(values)
            return [dict(zip(header, row)) for row in values if any(v is not None for v in row)]
        venues = rows("Заведения")
        reviews = rows("Отзывы")
        mentions = rows("Упоминания вкусов")
        workbook.close()
    verified = [r for r in reviews if r["source_status"] == "verified" and not r["duplicate_of"]]
    review_map = {r["ID отзыва"]: r for r in verified}
    venue_map = {r["venue_id"]: r for r in venues}
    concrete = [m for m in mentions if m["source_status"] == "verified"
                and m["ID отзыва"] in review_map and m["specificity"] == "concrete"]
    strict = [m for m in concrete if m["confidence"] == "high"
              and venue_map[m["venue_id"]]["sample_current"] == "yes"
              and review_map[m["ID отзыва"]]["platform_verification"] != "not_passed"
              and not review_map[m["ID отзыва"]].get("bias_flag")]
    return {
        "venues_total": len(venues), "current_sample": sum(v["sample_current"] == "yes" for v in venues),
        "reviews_total": len(reviews), "reviews_verified": len(verified),
        "mentions_total": len(mentions), "concrete_mentions": len(concrete),
        "concrete_reviews": len({m["ID отзыва"] for m in concrete}),
        "flavors": len({m["normalized_flavor"] for m in concrete}),
        "strict_mentions": len(strict), "strict_reviews": len({m["ID отзыва"] for m in strict}),
    }


def economics_summary():
    """Итоги готовых публичных срезов; не пересчёт закрытых рецептур."""
    menu = read_csv("data/menu_economics.csv")
    purchases = read_csv("data/menu_raw_material_costs.csv")
    used = sum((number(r["used_cost_rub"]) for r in menu), D(0))
    cash = sum((number(r["Закупка, ₽"]) for r in purchases), D(0))
    revenue = sum((number(r["price_50ml_rub"]) * number(r["portions"]) for r in menu), D(0))
    return {"ready_liters": sum(number(r["target_ready_ml"]) for r in menu) / 1000,
            "portions": sum(number(r["portions"]) for r in menu), "purchase_rub": cash,
            "used_rub": used, "leftover_rub": cash-used, "revenue_rub": revenue,
            "material_contribution_rub": revenue-used}


def build_artifacts():
    os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "bar-menu-mpl"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    menu = read_csv("data/menu_economics.csv")
    ink, gold, paper = "#21353F", "#9A6C43", "#FAF7F0"
    descriptions = {
        "C01": "Хрен · имбирь · мёд", "C02": "Вишня · ягодная классика",
        "C03": "Малина · ягодная коллекция", "C04": "Лимон · цедра и сок",
        "C05": "Щавель · мёд", "C09": "Зелёное яблоко · базилик",
        "C13": "Облепиха · ягодная коллекция", "C14": "Клюква · ягодная классика",
        "C06": "Подсолнечная халва · мёд", "C12": "Малина · молоко · сливки",
    }
    for directory in (ROOT / "results",):
        directory.mkdir(exist_ok=True)
    with plt.rc_context({"font.family": "DejaVu Sans", "pdf.fonttype": 42}):
        fig = plt.figure(figsize=(210/25.4, 297/25.4), facecolor=paper)
        ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
        ax.add_patch(plt.Rectangle((.065, .055), .87, .89, fill=False, edgecolor=gold, linewidth=.7))
        ax.text(.5, .901, "С А Н К Т - П Е Т Е Р Б У Р Г", ha="center", fontsize=8, color=gold)
        ax.text(.5, .842, "НАСТОЙКИ", ha="center", fontsize=33, fontfamily="DejaVu Serif", color=ink)
        ax.text(.5, .805, "Авторская коллекция", ha="center", fontsize=12, fontfamily="DejaVu Serif", style="italic", color=ink)
        ax.text(.5, .776, "50 мл · порция", ha="center", fontsize=9, color=gold)
        ax.plot([.44, .56], [.754, .754], color=gold, linewidth=.8)
        for tier, label, heading, start in [("basic", "ОСНОВНАЯ КОЛЛЕКЦИЯ", .722, .685),
                                            ("premium", "ДЕСЕРТНАЯ КОЛЛЕКЦИЯ", .264, .222)]:
            ax.text(.115, heading, label, fontsize=8.5, weight="bold", color=gold)
            for index, row in enumerate(r for r in menu if r["tier"] == tier):
                y = start-index*(.051 if tier == "basic" else .061)
                ax.text(.115, y, row["normalized_flavor"], fontsize=13.2, fontfamily="DejaVu Serif", color=ink)
                ax.text(.885, y, fmt(row["price_50ml_rub"], 0)+" ₽", fontsize=11.5, ha="right", weight="bold", color=ink)
                ax.text(.115, y-.018, descriptions[row["candidate_id"]], fontsize=7.6, color="#66736F")
        ax.text(.5, .092, "Малина — сливки содержит молочные продукты", ha="center", fontsize=7, color="#66736F")
        ax.set(xlim=(0, 1), ylim=(0, 1))
        fig.savefig(ROOT / ARTIFACTS[0], dpi=180, facecolor=paper)
        fig.savefig(ROOT / ARTIFACTS[2], facecolor=paper,
                    metadata={"Title": "Настойки — авторское меню", "Author": "Bar_menu",
                              "CreationDate": None, "ModDate": None})
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(10, 6), facecolor=paper)
        values = [float(number(r["cost_50ml_rub"])) for r in menu]
        ax.barh([r["normalized_flavor"] for r in menu], values,
                color=[gold if r["tier"] == "premium" else ink for r in menu], height=.65)
        for i, row in enumerate(menu):
            ax.text(values[i]+.4, i, fmt(row["cost_50ml_rub"])+" ₽", va="center", fontsize=9, color=ink)
        ax.set(title="Сырьё одной порции · 50 мл", xlabel="Рубли", xlim=(0, max(values)*1.23))
        ax.invert_yaxis(); ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="x", alpha=.1); ax.set_axisbelow(True)
        ax.legend(handles=[Line2D([0], [0], color=ink, lw=7, label="Основные"),
                           Line2D([0], [0], color=gold, lw=7, label="Премиальные")],
                  frameon=False, loc="lower right", fontsize=9)
        fig.text(.02, .025, "Ценовой сценарий 16.10.2026 · оценочные выходы · расчёт по готовой выгрузке", fontsize=8, color="#66736F")
        fig.tight_layout(rect=[0, .06, 1, 1]); fig.savefig(ROOT / ARTIFACTS[1], dpi=160); plt.close(fig)


def block_external_network():
    def audit(event, args):
        if event not in {"socket.connect", "socket.getaddrinfo"}:
            return
        if event == "socket.connect":
            if args[0].family == socket.AF_UNIX:
                return
            host = args[1][0]
        else:
            host = args[0]
        if host in {None, "localhost"}:
            return
        try:
            if ipaddress.ip_address(host).is_loopback:
                return
        except ValueError:
            pass
        raise RuntimeError("Публичный Notebook не должен обращаться к внешней сети")
    sys.addaudithook(audit)


def check_public_notebook(save=False):
    import nbformat
    from nbclient import NotebookClient
    before = {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ARTIFACTS}
    block_external_network(); build_artifacts()
    notebook = nbformat.read(ROOT / "bar_menu_analysis.ipynb", as_version=4)
    cells = [c for c in notebook.cells if c.cell_type == "code"]
    for cell in cells:
        cell.outputs = []; cell.execution_count = None
    original = cells[0].source
    cells[0].source = "from scripts.portfolio import block_external_network\nblock_external_network()\n" + original
    NotebookClient(notebook, timeout=180, kernel_name="python3",
                   resources={"metadata": {"path": str(ROOT)}}).execute()
    cells[0].source = original
    if save:
        nbformat.write(notebook, ROOT / "bar_menu_analysis.ipynb")
    after = {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ARTIFACTS}
    if before != after:
        raise RuntimeError("PDF/PNG отличаются от сохранённых результатов")
    print(f"Public notebook: {len(cells)} cells, fresh kernel, no errors; PDF/PNG reproduced")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--save-notebook", action="store_true")
    args = parser.parse_args()
    if args.check or args.save_notebook:
        check_public_notebook(save=args.save_notebook)
    else:
        build_artifacts()
