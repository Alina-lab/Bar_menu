"""Проверка публичных выгрузок и сборка меню; исходных рецептур здесь нет."""
import argparse
import csv
import hashlib
import ipaddress
import os
import socket
import sys
import tempfile
from decimal import Decimal, ROUND_HALF_UP, getcontext
from pathlib import Path

getcontext().prec = 50
PROJECT_ROOT = Path(__file__).resolve().parents[2]
ROOT = PROJECT_ROOT / "technical"
D = Decimal
ARTIFACTS = ("results/menu_preview.png", "results/cost_50ml.png", "results/nastoyka_menu.pdf")
PREMIUM = {"C06", "C12"}
INK, ACCENT, PAPER = "#3D303A", "#7A805E", "#FCF9F2"
MENU_GROUPS = (
    ("Сладкие и десертные", ("C02", "C03", "C06", "C12")),
    ("Кислые и освежающие", ("C04", "C05", "C13", "C14")),
    ("Свежие и травяные", ("C09",)),
    ("Острые и пряные", ("C01",)),
)
MENU_DESCRIPTIONS = {
    "C01": "Корень хрена, имбирь и мёд",
    "C02": "Вишня без косточек",
    "C03": "Малиновая настойка",
    "C04": "Лимонная цедра и сок",
    "C05": "Щавель с мёдом",
    "C09": "Зелёное яблоко и базилик",
    "C13": "Настойка на облепихе",
    "C14": "Ягоды клюквы",
    "C06": "Подсолнечная халва и мёд",
    "C12": "Малина, молоко и сливки",
}


def read_csv(name):
    with (ROOT / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def number(value):
    return D(str(value).replace("\u00a0", "").replace(" ", "").replace(",", "."))


def fmt(value, digits=2):
    value = number(value).quantize(D(1).scaleb(-digits), rounding=ROUND_HALF_UP)
    return f"{value:,.{digits}f}".replace(",", " ").replace(".", ",")


def research_summary():
    """Повторить прежние определения срезов из исходного Excel, не меняя книгу."""
    from openpyxl import load_workbook
    with (ROOT / "data/nastoyki_reviews.xlsx").open("rb") as stream:
        workbook = load_workbook(stream, read_only=True, data_only=True)
        def rows(sheet):
            values = workbook[sheet].iter_rows(values_only=True)
            header = next(values)
            return [dict(zip(header, row)) for row in values if any(v is not None for v in row)]
        venues = rows("Заведения")
        reviews = rows("Отзывы")
        mentions = rows("Упоминания вкусов")
        workbook.close()
    verified = [r for r in reviews if r["source_status"] == "verified" and not r["duplicate_of"]]
    review_map = {r["ID отзыва"]: r for r in verified}
    venue_map = {r["venue_id"]: r for r in venues}
    concrete = [m for m in mentions if m["source_status"] == "verified"
                and m["ID отзыва"] in review_map and m["specificity"] == "concrete"]
    strict = [m for m in concrete if m["confidence"] == "high"
              and venue_map[m["venue_id"]]["sample_current"] == "yes"
              and review_map[m["ID отзыва"]]["platform_verification"] != "not_passed"
              and not review_map[m["ID отзыва"]].get("bias_flag")]
    return {
        "venues_total": len(venues), "current_sample": sum(v["sample_current"] == "yes" for v in venues),
        "reviews_total": len(reviews), "reviews_verified": len(verified),
        "mentions_total": len(mentions), "concrete_mentions": len(concrete),
        "concrete_reviews": len({m["ID отзыва"] for m in concrete}),
        "flavors": len({m["normalized_flavor"] for m in concrete}),
        "strict_mentions": len(strict), "strict_reviews": len({m["ID отзыва"] for m in strict}),
    }


def economics_summary():
    """Итоги готовых публичных срезов; не пересчёт закрытых рецептур."""
    menu = read_csv("data/menu_economics.csv")
    purchases = read_csv("data/menu_raw_material_costs.csv")
    used = sum((number(r["used_cost_rub"]) for r in menu), D(0))
    cash = sum((number(r["Закупка, ₽"]) for r in purchases), D(0))
    revenue = sum((number(r["price_50ml_rub"]) * number(r["portions"]) for r in menu), D(0))
    return {"ready_liters": sum(number(r["target_ready_ml"]) for r in menu) / 1000,
            "portions": sum(number(r["portions"]) for r in menu), "purchase_rub": cash,
            "used_rub": used, "leftover_rub": cash-used, "revenue_rub": revenue,
            "material_contribution_rub": revenue-used}


def build_artifacts():
    os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "bar-menu-mpl"))
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D

    menu = read_csv("data/menu_economics.csv")
    by_id = {row["candidate_id"]: row for row in menu}
    ink, accent, paper = INK, ACCENT, PAPER
    (ROOT / "results").mkdir(exist_ok=True)
    with plt.rc_context({"font.family": "DejaVu Sans", "pdf.fonttype": 42}):
        from matplotlib.path import Path as VectorPath
        from matplotlib.patches import PathPatch

        menu_ink, menu_accent, menu_paper = "#352E2A", "#825445", "#FAF6EC"
        page_width, page_height = 210/25.4*72, 297/25.4*72
        fig = plt.figure(figsize=(page_width/72, page_height/72), facecolor=menu_paper)
        ax = fig.add_axes([0, 0, 1, 1]); ax.axis("off")
        ax.set(xlim=(0, page_width), ylim=(0, page_height))
        left, names_x, right = 48, 174, page_width-48
        ax.text(left, 766, "НАСТОЙКИ", fontsize=33, fontfamily="DejaVu Serif", color=menu_ink)
        ax.text(right, 768, "50 мл", ha="right", fontsize=10.5, color=menu_accent)
        ax.plot([left, right], [737, 737], color=menu_accent, linewidth=.7)

        # One continuous rhythm for all ten drinks; category labels use the side margin.
        y, row_step = 691, 60
        for index, (label, ids) in enumerate(MENU_GROUPS):
            if index:
                ax.plot([left, right], [y+27, y+27], color="#D8CFC2", linewidth=.45)
            category = label.upper().replace(" И ", "\nИ ")
            ax.text(left, y+1, category, fontsize=8.5, color=menu_accent,
                    va="top", linespacing=1.55)
            for cid in ids:
                row = by_id[cid]
                ax.text(names_x, y, row["normalized_flavor"], fontsize=15.2,
                        fontfamily="DejaVu Serif", color=menu_ink)
                ax.text(names_x, y-17, MENU_DESCRIPTIONS[cid], fontsize=9.2, color="#71665C")
                ax.text(right, y, fmt(row["price_50ml_rub"], 0)+" ₽", fontsize=12.1,
                        ha="right", color=menu_ink)
                y -= row_step

        # A small open branch, without leaves, fruit or filled decorative shapes.
        centre = page_width/2
        for points in [
            [(centre-27, 72), (centre-8, 81), (centre+8, 67), (centre+27, 76)],
            [(centre-10, 75), (centre-13, 79), (centre-13, 82), (centre-11, 85)],
            [(centre+10, 72), (centre+13, 70), (centre+16, 67), (centre+18, 68)],
        ]:
            path = VectorPath(points, [VectorPath.MOVETO]+[VectorPath.CURVE4]*3)
            ax.add_patch(PathPatch(path, fill=False, edgecolor="#AC9580", linewidth=.55,
                                   capstyle="round"))
        fig.savefig(ROOT / ARTIFACTS[0], dpi=180, facecolor=menu_paper)
        fig.savefig(ROOT / ARTIFACTS[2], facecolor=menu_paper,
                    metadata={"Title": "Настойки — барное меню", "Author": "Bar_menu",
                              "CreationDate": None, "ModDate": None})
        plt.close(fig)
        fig, ax = plt.subplots(figsize=(10, 6), facecolor=paper)
        values = [float(number(r["cost_50ml_rub"])) for r in menu]
        ax.barh([r["normalized_flavor"] for r in menu], values,
                color=[accent if r["tier"] == "premium" else ink for r in menu], height=.65)
        for i, row in enumerate(menu):
            ax.text(values[i]+.4, i, fmt(row["cost_50ml_rub"])+" ₽", va="center", fontsize=9, color=ink)
        ax.set(title="Сырьё одной порции · 50 мл", xlabel="Рубли", xlim=(0, max(values)*1.23))
        ax.invert_yaxis(); ax.spines[["top", "right"]].set_visible(False)
        ax.grid(axis="x", alpha=.1); ax.set_axisbelow(True)
        ax.legend(handles=[Line2D([0], [0], color=ink, lw=7, label="Основные"),
                           Line2D([0], [0], color=accent, lw=7, label="Премиальные")],
                  frameon=False, loc="lower right", fontsize=9)
        fig.text(.02, .025, "Ценовой сценарий 16.10.2026 · оценочные выходы · расчёт по готовой выгрузке", fontsize=8, color="#66736F")
        fig.tight_layout(rect=[0, .06, 1, 1]); fig.savefig(ROOT / ARTIFACTS[1], dpi=160); plt.close(fig)


def block_external_network():
    def audit(event, args):
        if event not in {"socket.connect", "socket.getaddrinfo"}:
            return
        if event == "socket.connect":
            if args[0].family == socket.AF_UNIX:
                return
            host = args[1][0]
        else:
            host = args[0]
        if host in {None, "localhost"}:
            return
        try:
            if ipaddress.ip_address(host).is_loopback:
                return
        except ValueError:
            pass
        raise RuntimeError("Публичный Notebook не должен обращаться к внешней сети")
    sys.addaudithook(audit)


def check_public_notebook(save=False):
    import nbformat
    from nbclient import NotebookClient
    before = {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ARTIFACTS}
    block_external_network(); build_artifacts()
    notebook = nbformat.read(PROJECT_ROOT / "bar_menu_analysis.ipynb", as_version=4)
    cells = [c for c in notebook.cells if c.cell_type == "code"]
    for cell in cells:
        cell.outputs = []; cell.execution_count = None
    original = cells[0].source
    cells[0].source = "from technical.scripts.portfolio import block_external_network\nblock_external_network()\n" + original
    NotebookClient(notebook, timeout=180, kernel_name="python3",
                   resources={"metadata": {"path": str(PROJECT_ROOT)}}).execute()
    cells[0].source = original
    if save:
        nbformat.write(notebook, PROJECT_ROOT / "bar_menu_analysis.ipynb")
    after = {p: hashlib.sha256((ROOT/p).read_bytes()).hexdigest() for p in ARTIFACTS}
    if before != after:
        raise RuntimeError("PDF/PNG отличаются от сохранённых результатов")
    print(f"Public notebook: {len(cells)} cells, fresh kernel, no errors; PDF/PNG reproduced")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--check", action="store_true")
    parser.add_argument("--save-notebook", action="store_true")
    args = parser.parse_args()
    if args.check or args.save_notebook:
        check_public_notebook(save=args.save_notebook)
    else:
        build_artifacts()
